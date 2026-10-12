# -*- coding: utf-8 -*-
"""
DOCX 笔记动态生成器 v2.1

解决 gen_p01_v1.py 频繁修改导致的 Edit 竞态问题。

使用方法：
1. 创建 JSON 配置文件（如 config_p01.json），包含 TITLE/SOURCE/OUTPUT_PATH/FRAMES_DIR/FRAMES/SUMMARY/SECTIONS
2. 运行: python gen_docx_dynamic.py --config config_p01.json

JSON 配置示例见 templates/config_example.json

SECTIONS 元素格式（与 docx_note_v2.py 相同）：
  ["h1", "一、xxx"]                    顶级大标题
  ["h2", "1. xxx"]                      二级标题
  ["h3", "① xxx"]                      三级标题
  ["body", "正文..."]                  普通正文
  ["red", "关键术语"]                  红色加粗
  ["bold", "重要标题"]                黑色加粗
  ["bullet", [[text, bold, red], ...]] 项目符号列表
  ["table", [headers...], [[row1...], [row2...]]] 表格
  ["img", "0001", "caption"]           图片
  ["formula", "C = λF"]                公式
  ["why", "解释..."]                    WHY 解释
  ["tips", "做题要点..."]              做题要点
"""
import os
import sys
import re
import json
import argparse
from pathlib import Path
from io import BytesIO
from PIL import Image

sys.path = [p for p in sys.path if 'hermes' not in p]

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
# Package imports and direct/importlib execution share a loader without changing sys.path.
if __package__:
    from .formula_layout import add_text as _add_formula_text
else:
    import importlib.util as _formula_util
    from pathlib import Path as _FormulaPath
    _formula_name = '_bili_formula_layout'
    if _formula_name not in sys.modules:
        _formula_root = _FormulaPath(__file__).resolve().parent / 'formula_layout'
        _formula_spec = _formula_util.spec_from_file_location(
            _formula_name, _formula_root / '__init__.py',
            submodule_search_locations=[str(_formula_root)])
        _formula_module = _formula_util.module_from_spec(_formula_spec)
        sys.modules[_formula_name] = _formula_module
        try:
            _formula_spec.loader.exec_module(_formula_module)
        except Exception:
            sys.modules.pop(_formula_name, None)
            raise
    _add_formula_text = sys.modules[_formula_name].add_text

# =============================================================================
# === 默认样式常量 ===
# =============================================================================

PAGE_MARGIN = Cm(1.27)
BODY_SIZE = 13
RED_COLOR = RGBColor(192, 0, 0)
H1_SIZE = 17
H2_SIZE = 15
H3_SIZE = 14
RED_SIZE = 14
FORMULA_SIZE = 15
LINE_SPACING_BODY = 1.4
SPACE_AFTER_P = Pt(2)


def set_run_font(run, name='宋体', size=BODY_SIZE, bold=False, red=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    if red:
        run.font.color.rgb = RED_COLOR


def _add_runs_with_sub(p, text, name='宋体', size=BODY_SIZE, bold=False, red=False):
    """统一排版正文、上下标与补事件横线；数学结构使用 Word OMML。"""
    _add_formula_text(p, text, name=name, size=size, bold=bold, red=red)


def setup_section(doc):
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = PAGE_MARGIN
    return s


def add_heading_styled(doc, text, level, size):
    style_name = f'Heading {level}'
    try:
        style = doc.styles[style_name]
    except KeyError:
        style = doc.styles.add_style(style_name, 1)
    p = doc.add_paragraph()
    p.style = style
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(8 if level == 1 else 6 if level == 2 else 4)
    p.paragraph_format.space_after = Pt(4 if level == 1 else 2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.2
    if p.runs:
        p._element.remove(p.runs[0]._element)
    _add_runs_with_sub(p, text, name='黑体', size=size, bold=True)
    return p


def add_h1(doc, text): return add_heading_styled(doc, text, 1, H1_SIZE)
def add_h2(doc, text): return add_heading_styled(doc, text, 2, H2_SIZE)
def add_h3(doc, text): return add_heading_styled(doc, text, 3, H3_SIZE)


def add_body(doc, text, bold=False, red=False, size=BODY_SIZE):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = LINE_SPACING_BODY
    _add_runs_with_sub(p, text, size=size, bold=bold, red=red)
    return p


def add_bullet(doc, parts, size=BODY_SIZE):
    """每个 parts 条目独立成一个 bullet 段落。"""
    for text, is_bold, is_red in parts:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = LINE_SPACING_BODY
        _add_runs_with_sub(p, text, size=size, bold=is_bold, red=is_red)
    return doc


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        p = table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_runs_with_sub(p, h, name='黑体', size=BODY_SIZE, bold=True)
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            p = table.rows[r_idx + 1].cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs_with_sub(p, cell_text, size=BODY_SIZE - 1)
    doc.add_paragraph()
    return table


def resolve_frame(frames_dir, frame_id):
    name = str(frame_id)
    if not name.endswith('.jpg'):
        name = (name if name.startswith('frame_') else 'frame_' + name) + '.jpg'
    root = Path(frames_dir).resolve(); path = (root/name).resolve()
    if path.parent != root or not path.is_file():
        raise ValueError(f'Missing or invalid frame: {name}')
    return path

def add_img(doc, frame_id, caption, img_width, frames_dir):
    path = resolve_frame(frames_dir, frame_id)
    # Normalize ffmpeg JPEG headers before python-docx reads the stream.
    data = BytesIO()
    with Image.open(path) as image:
        image.convert('RGB').save(data,format='JPEG',quality=95)
    data.seek(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(data,width=img_width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.add_run(f'图：{caption}'),size=11)
    return p


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    _add_runs_with_sub(p, text, name='Cambria Math', size=FORMULA_SIZE, bold=True, red=True)
    return p


def add_why(doc, text):
    label = add_body(doc, 'WHY 解释：', bold=True, red=True, size=RED_SIZE)
    label.paragraph_format.keep_with_next = True
    add_body(doc, text)


def add_tips(doc, text):
    label = add_body(doc, '做题要点', bold=True, red=True, size=RED_SIZE)
    label.paragraph_format.keep_with_next = True
    add_bullet(doc, [(text, False, False)])


# =============================================================================
# === 主函数 ===
# =============================================================================

def build_docx(config):
    """根据配置字典生成 DOCX。"""
    doc = Document()
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    section = setup_section(doc)
    img_width = section.page_width - section.left_margin - section.right_margin

    title = config['TITLE']
    source = config['SOURCE']
    output_path = config['OUTPUT_PATH']
    frames_dir = config['FRAMES_DIR']
    frames = config.get('FRAMES', {})
    summary = config.get('SUMMARY', [])
    sections = config.get('SECTIONS', [])

    # 标题
    if re.match(r'^第[\u4e00-\u9fa5]+章', title):
        add_heading_styled(doc, title, 1, H1_SIZE + 3)
    elif re.match(r'^\d+\.\d+\.\d+', title):
        add_heading_styled(doc, title, 3, H2_SIZE)
    elif re.match(r'^\d+\.\d+', title):
        add_heading_styled(doc, title, 2, H1_SIZE)
    else:
        add_heading_styled(doc, title, 1, H1_SIZE + 3)

    # Keep the document title independent of the heading hierarchy.
    title_paragraph = doc.paragraphs[0]
    title_paragraph.style = doc.styles['Title']
    title_paragraph.style.font.color.rgb = RGBColor(0, 0, 0)
    for run in title_paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    # Word's default styles may carry a decorative paragraph border.
    for paragraph_style in doc.styles:
        for border in paragraph_style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for paragraph in doc.paragraphs:
        for border in paragraph._p.xpath('.//w:pBdr'):
            border.getparent().remove(border)

    # 来源
    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t2.paragraph_format.space_after = Pt(4)
    source_label = source if source.lstrip().startswith(('来源：', '来源:')) else f'来源：{source}'
    set_run_font(t2.add_run(source_label), size=11)
    doc.add_paragraph()

    # 内容区
    for item in sections:
        typ = item[0]
        if typ == 'h1':
            add_h1(doc, item[1])
        elif typ == 'h2':
            add_h2(doc, item[1])
        elif typ == 'h3':
            add_h3(doc, item[1])
        elif typ == 'body':
            add_body(doc, item[1])
        elif typ == 'red':
            add_body(doc, item[1], bold=True, red=True, size=RED_SIZE)
        elif typ == 'bold':
            add_body(doc, item[1], bold=True)
        elif typ == 'bullet':
            add_bullet(doc, item[1])
        elif typ == 'table':
            add_table(doc, item[1], item[2])
        elif typ == 'img':
            add_img(doc, item[1], item[2], img_width, frames_dir)
        elif typ == 'formula':
            add_formula(doc, item[1])
        elif typ == 'why':
            add_why(doc, item[1])
        elif typ == 'tips':
            add_tips(doc, item[1])

    # 总结
    add_h1(doc, '本节要点总结')
    for s in summary:
        add_bullet(doc, [(s, False, False)])

    doc.core_properties.title = title
    doc.core_properties.subject = source
    doc.core_properties.author = 'AI Agent'
    doc.save(output_path)

    # 验证 markdown 残留
    import zipfile
    with zipfile.ZipFile(output_path) as z:
        with z.open('word/document.xml') as f:
            xml = f.read().decode('utf-8')
    n = len(re.findall(r'\*\*[^*]+\*\*', xml))
    print(f'Saved: {output_path}\nSize: {os.path.getsize(output_path)/1024:.1f} KB')
    print('No markdown residue.' if n == 0 else f'WARNING: {n} markdown residues!')

    return output_path


def main():
    parser = argparse.ArgumentParser(description='DOCX 笔记动态生成器')
    parser.add_argument('--config', '-c', required=True, help='JSON 配置文件路径')
    parser.add_argument('--sample-output', help='生成两张截图和首张表格的样本文档；不覆盖正式输出')
    args = parser.parse_args()

    with open(args.config, 'r', encoding='utf-8') as f:
        config = json.load(f)

    # Validate every referenced image before writing any document.
    for item in config.get('SECTIONS',[]):
        if item[0] == 'img':
            with Image.open(resolve_frame(config['FRAMES_DIR'],item[1])) as image:
                image.load()
    if args.sample_output:
        chosen=[]; counts={'img':0,'table':0}
        for item in config.get('SECTIONS',[]):
            kind=item[0]
            if kind in counts and counts[kind] < (2 if kind=='img' else 1):
                chosen.append(item);counts[kind]+=1
        config={**config,'OUTPUT_PATH':args.sample_output,'SECTIONS':[['h1','样本排版检查'],['body','检查字体、表格、截图与图注后，再生成全文。']]+chosen}
    build_docx(config)


if __name__ == '__main__':
    main()
