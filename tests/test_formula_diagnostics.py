"""New step-2 acceptance tests; archived parser trees are not the oracle."""
import json
import sys
import unittest
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'templates'))
from formula_layout import parse, add_text
from formula_layout.operators import OPERATORS, NOTICE_CODES

FIXTURES=Path(__file__).parent/'fixtures/generator_baseline'
def cases(): return json.loads((FIXTURES/'formula_cases.json').read_text(encoding='utf-8'))['cases']
def walk(nodes):
    for n in nodes:
        yield n
        yield from walk(n.children)
def text(nodes):
    out=''
    for n in nodes:
        if n.kind=='text':out+=n.text
        elif n.kind=='group':out+=n.text[0]+text(n.children)+n.text[1]
        elif n.kind in ('sup','sub'):out+=text(n.children[:1])+('^' if n.kind=='sup' else '_')+'{'+text(n.children[1].children)+'}'
        elif n.kind=='nary':out+=n.text
        else:out+=text(n.children)
    return out

class UnsupportedAcceptance(unittest.TestCase):
    def test_all_nary_cases_recognized_in_generic_domain(self):
        for c in cases():
            if c['classification']!='nary': continue
            with self.subTest(case=c['id']):
                r=parse(c['input'],'generic'); nary=[n for n in walk(r.nodes) if n.kind=='nary']
                self.assertEqual(len(nary),c.get('nary_count',1))
                n=nary[0];self.assertEqual(text(n.lower or ()),c['lower'] or '')
                self.assertEqual(text(n.upper or ()),c['upper'] or '')
                self.assertEqual(text(n.children),text(parse(c['body'],'generic').nodes))
                self.assertTrue(n.text.endswith('{'+c['body']+'}'))
                self.assertEqual(len(r.notices),len(nary))
                self.assertTrue(all(x.code==c['code'] and x.severity==c['step2_severity'] for x in r.notices))

    def test_legacy_operators_warn_without_blocking(self):
        for c in cases():
            if c['classification']=='legacy_nary':
                r=parse(c['input']);self.assertEqual(r.notices[0].code,c['code'])
                self.assertEqual(r.notices[0].severity,'warning')

    def test_sum_aliases_have_equal_recognition(self):
        for op in ('Σ','∑'):
            r=parse(op+'_{k=0}^{n}{p_k}')
            self.assertEqual(r.nodes[0].kind,'nary')
            self.assertEqual(OPERATORS[op].chr,'∑')
            self.assertEqual(OPERATORS[op].limLoc,'undOvr')
            self.assertEqual(r.notices[0].code,'UNSUPPORTED_NARY')

    def test_warning_operators_are_literal_without_new_omml_emission(self):
        for c in cases():
            if c['classification']!='nary':continue
            d=Document();p=d.add_paragraph();notices=add_text(p,c['input'],domain='generic')
            if c['operator']!='Σ': self.assertEqual(p.text,c['input'])
            else:
                self.assertTrue(list(p._p.iter(qn('m:sSub'))))
                self.assertTrue(list(p._p.iter(qn('m:sSup'))))
            self.assertFalse(list(p._p.iter(qn('m:nary'))))
            self.assertTrue(notices)

    def test_unknown_domain_uses_notice_channel(self):
        r=parse('x^2','absent');self.assertEqual(r.notices[0].code,'UNKNOWN_DOMAIN')
        self.assertIn(r.notices[0].code,NOTICE_CODES)

    def test_legacy_infinite_and_greek_limits_remain_warnings(self):
        for s in ('∫_0^∞ f(x) dx','Σ_α^β f(x)'):
            r=parse(s);self.assertEqual([(n.code,n.severity) for n in r.notices],[('LEGACY_NARY_BODY','warning')])

class SilentlyWrongAcceptance(unittest.TestCase):
    def check_case(self,name):
        c=next(c for c in cases() if c['id']==name);r=parse(c['input'])
        self.assertEqual([(n.code,n.severity) for n in r.notices],[(c['code'],'error')])
        self.assertEqual(r.nodes[0].text,c['input'])
        self.assertEqual(r.notices[0].text_span,(0,len(c['input'])))
        d=Document();p=d.add_paragraph();returned=add_text(p,c['input'])
        self.assertEqual(returned[0].code,c['code'])
        self.assertEqual(p.text,c['input'])
        self.assertFalse(list(p._p.iter(qn('m:sSup'))))
        self.assertFalse(list(p._p.iter(qn('m:sSub'))))
        self.assertFalse(list(p._p.iter(qn('m:bar'))))
    def test_empty_sup_is_error_not_empty_node(self):self.check_case('empty_sup')
    def test_empty_sub_is_error_not_empty_node(self):self.check_case('empty_sub')
    def test_cost_is_error_not_complement_plus_ost(self):self.check_case('ambiguous_cost')
    def test_power_does_not_swallow_following_variable(self):self.check_case('ambiguous_power')
    def test_explicit_disambiguation_preserves_valid_math(self):
        for s in ('x^{2}y','x^{2y}','A^{c}ost','A^cB^c','A_{ij}^c','p^k(1−p)^(n−k)'):
            self.assertFalse(parse(s).notices,s)
    def test_no_strict_exception_channel_remains(self):
        import inspect
        self.assertNotIn('strict',inspect.signature(parse).parameters)
        self.assertNotIn('strict',inspect.signature(add_text).parameters)

class DetectedAcceptance(unittest.TestCase):
    def test_warnings_collected_before_error_are_retained_in_order(self):
        r=parse('∫_{0}^{1}{1} dx + A^cost')
        self.assertEqual([(n.code,n.severity) for n in r.notices],
            [('UNSUPPORTED_NARY','warning'),('AMBIGUOUS_SCRIPT','error')])
        self.assertEqual(r.notices[0].text_span,(0,12))
        self.assertEqual(r.notices[1].text_span,(18,24))
        self.assertEqual(r.nodes[0].kind,'text')

    def test_legacy_warning_survives_following_error_and_reaches_add_text(self):
        d=Document();p=d.add_paragraph()
        n=add_text(p,'∫_0^1 dy + A_{}',block_id='mixed')
        self.assertEqual([(x.code,x.severity) for x in n],[('LEGACY_NARY_BODY','warning'),('EMPTY_SCRIPT','error')])
        self.assertTrue(all(x.block_id=='mixed' for x in n))

    def test_plain_prose_caret_boundary_is_explicit(self):
        for s,code in [('正文 ^ 提示','MISSING_OPERAND'),('成本^收益','INVALID_SCRIPT'),('^','MISSING_OPERAND')]:
            self.assertEqual([(n.code,n.severity) for n in parse(s).notices],[(code,'error')])
        self.assertFalse(parse('正文用文字描述脱字符').notices)
        self.assertFalse(parse('x^n').notices)

    def test_continuous_complements_and_multichar_exponents(self):
        for s in ('A^cB^c','A^{c}b','x^{2y}','x^{cost}','x^12'):
            self.assertFalse(parse(s,'generic' if s=='x^{cost}' else 'mathematics').notices,s)
        for s in ('A^cfoo','x^2y','x^nn'):
            self.assertEqual(parse(s).notices[0].code,'AMBIGUOUS_SCRIPT')

    def test_actual_template_docx_keeps_legacy_sigma_layout(self):
        import importlib.util,tempfile
        spec=importlib.util.spec_from_file_location('sigma_regression',Path(__file__).resolve().parents[1]/'templates/gen_docx_dynamic.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'sum.docx'
            module.build_docx({'TITLE':'固定求和测试','SOURCE':'合成素材','OUTPUT_PATH':str(out),'FRAMES_DIR':tmp,'SUMMARY':[], 'SECTIONS':[['formula','Σ_{k=0}^{n}{p_k}']]})
            d=Document(out);sup=next(d.element.iter(qn('m:sSup')))
            sub=sup.find(qn('m:e')).find(qn('m:sSub'))
            self.assertIsNotNone(sub)
            def content(el):return ''.join(n.text or '' for n in el.iter(qn('m:t')))
            self.assertEqual(content(sub.find(qn('m:e'))),'Σ')
            self.assertEqual(content(sub.find(qn('m:sub'))),'k=0')
            self.assertEqual(content(sup.find(qn('m:sup'))),'n')
            self.assertFalse(list(d.element.iter(qn('m:nary'))))
            notice=add_text(d.add_paragraph(),'Σ_{k=0}^{n}{p_k}')[0]
            self.assertEqual((notice.code,notice.severity),('UNSUPPORTED_NARY','warning'))

    def test_sigma_without_braced_body_keeps_limits_without_inventing_braces(self):
        d=Document();p=d.add_paragraph()
        notices=add_text(p,'Σ_{k=0}^{n}p_k')
        self.assertEqual(notices[0].code,'LEGACY_NARY_BODY')
        self.assertEqual(len(list(p._p.iter(qn('m:sSup')))),1)
        self.assertEqual(len(list(p._p.iter(qn('m:sSub')))),2)
        content=''.join(n.text or '' for n in p._p.iter(qn('m:t')))
        self.assertEqual(content,'Σk=0npk')

    def test_six_previously_detected_cases_have_typed_errors(self):
        snapshots=json.loads((FIXTURES/'current_parser_snapshot.json').read_text(encoding='utf-8'))['cases']
        ids={c['id'] for c in snapshots if c['current_state']=='detected'}
        self.assertEqual(len(ids),6)
        for c in cases():
            if c['id'] in ids:
                with self.subTest(case=c['id']):
                    r=parse(c['input']);self.assertEqual([(n.code,n.severity) for n in r.notices],[(c['code'],'error')])
                    self.assertIn(r.notices[0].code,NOTICE_CODES)

    def test_nary_warning_never_masks_malformed_body(self):
        for s in ('∫_{0}^{1}{A^cost} dx','Σ_{0}^{n}{A_{} }','∬_{D}{A_{ij}'):
            r=parse(s);self.assertTrue(r.notices)
            self.assertTrue(all(n.severity=='error' for n in r.notices))

    def test_coordinates_match_fixed_examples(self):
        from formula_layout.notices import diagnose_runs
        for c in json.loads((FIXTURES/'notice_locations.json').read_text(encoding='utf-8'))['cases']:
            notices=diagnose_runs(c['text'],c['block_id'],c['sub_path'])
            self.assertEqual(notices[0].code,c['code'])
            self.assertEqual(list(notices[0].text_span),c['text_span'])
            self.assertEqual(notices[0].sub_path,c['sub_path'])

    def test_whitespace_or_empty_nested_groups_are_not_script_arguments(self):
        for s,code in [('A^{ }','EMPTY_SCRIPT'),('A^{()}','EMPTY_SCRIPT'),('∫_{0}^{1}{()} dx','EMPTY_NARY_BODY')]:
            self.assertEqual(parse(s).notices[0].code,code)

    def test_all_23_specs_have_expected_step2_outcomes(self):
        for c in cases():
            with self.subTest(case=c['id']):
                r=parse(c['input'],c.get('domain','mathematics'))
                if c['classification']=='supported':
                    self.assertFalse(r.notices)
                    counts={k:sum(n.kind==k for n in walk(r.nodes)) for k in c['expected_kind_counts']}
                    self.assertEqual(counts,c['expected_kind_counts'])
                else:
                    self.assertTrue(r.notices)
                    self.assertTrue(all((n.code,n.severity)==(c['code'],c['step2_severity']) for n in r.notices))

    def test_public_notice_shape_and_offset_mapping(self):
        from dataclasses import asdict
        n=add_text(Document().add_paragraph(),'概率 A_{ij',block_id='b',sub_path={'row':2,'col':1},text_offset=4)[0]
        self.assertEqual(set(asdict(n)),{'code','severity','block_id','sub_path','text_span','message'})
        self.assertEqual(n.text_span,(7,12))

    def test_duplicate_limits_and_stray_closing_delimiters_are_errors(self):
        self.assertEqual(parse('∫_{0}_{1}{x}').notices[0].code,'DUPLICATE_LIMIT')
        self.assertEqual(parse('A^{n]}').notices[0].code,'UNEXPECTED_CLOSING')

    def test_merged_runs_use_unicode_offsets_and_render_row_numbers(self):
        from formula_layout.notices import diagnose_runs
        n = (diagnose_runs([{'text':'A^'},{'text':'{}'}], 'legacy-1', {'row':0,'col':0})
             + diagnose_runs([{'text':'概率 A_'},{'text':'{ij'}], 'legacy-1', {'row':1,'col':0})
             + diagnose_runs([{'text':'🙂 A^'},{'text':'cost'}], 'legacy-2', {'item':0}))
        self.assertEqual([(x.block_id,x.sub_path,x.text_span) for x in n],
            [('legacy-1',{'row':0,'col':0},(0,4)),('legacy-1',{'row':1,'col':0},(3,8)),('legacy-2',{'item':0},(2,8))])

if __name__=='__main__':unittest.main()
