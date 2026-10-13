"""Regression checks for notation semantics and all generator entry points."""
import sys
import importlib.util
from pathlib import Path
import tempfile
import unittest
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'templates'))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from formula_layout import add_text, parse, register_domain

class FormulaTests(unittest.TestCase):
    def test_complements(self):
        d=Document(); p=d.add_paragraph()
        add_text(p,'全不发生 A^cB^{c}C^c；[A(B^c∪C)D]^c',size=11,bold=True,red=True)
        bars=list(p._p.iter(qn('m:bar')))
        self.assertEqual(len(bars),5)
        self.assertTrue(any(len(list(bar.iter(qn('m:bar'))))==2 for bar in bars))
        self.assertNotIn('^c',p._p.xml)
        self.assertIn('C00000',p._p.xml)
        self.assertIn('\u2009',p._p.xml)
    def test_generic_and_scripts(self):
        self.assertEqual(parse('A^c','generic').nodes[-1].kind,'sup')
        d=Document();p=d.add_paragraph();add_text(p,'A_{ij}^c 与 x^{n+1}')
        self.assertEqual(len(list(p._p.iter(qn('m:sSub')))),1)
        self.assertEqual(len(list(p._p.iter(qn('m:sSup')))),1)
        self.assertEqual(len(list(p._p.iter(qn('m:bar')))),1)
        expression=parse('p^k(1−p)^(n−k)').nodes
        self.assertEqual(sum(n.kind=='sup' for n in expression),2)
    def test_domain_registration(self):
        register_domain('test_physics',lambda op,arg:('bar',3) if op=='^' and arg.startswith('avg') else None)
        self.assertEqual(parse('v^{avg}','test_physics').nodes[-1].kind,'bar')
    def test_invalid_and_prose(self):
        self.assertEqual(parse('正文（普通文字）').nodes[0].text,'正文（普通文字）')
        for s in ('^c','A^','(AB^c','A_'):
            self.assertTrue(any(n.severity == 'error' for n in parse(s).notices))
    def test_generators(self):
        for filename in ('gen_docx_dynamic.py','docx_note_v2.py'):
            path=Path(__file__).resolve().parents[1]/'templates'/filename
            spec=importlib.util.spec_from_file_location('test_generator',path)
            module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
            d=Document();module.add_body(d,'正文 A^c');module.add_formula(d,'(AB)^c')
            module.add_table(d,['事件'],[['A^cB^c']])
            self.assertEqual(len(list(d.element.iter(qn('m:bar')))),4)
            with tempfile.TemporaryDirectory() as tmp:
                path=Path(tmp)/'test.docx';d.save(path);reloaded=Document(path)
                self.assertEqual(len(list(reloaded.element.iter(qn('m:bar')))),4)
    def test_verifier_rejects_split_plain_notation(self):
        path=Path(__file__).resolve().parents[1]/'scripts'/'verify_docx.py'
        spec=importlib.util.spec_from_file_location('verify_test',path);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as tmp:
            file=Path(tmp)/'bad.docx';d=Document();p=d.add_paragraph();p.add_run('A^');p.add_run('c');d.save(file)
            self.assertFalse(module.verify(str(file),[],min_images=0))

    def test_prose_fallback(self):
        for text in ('填空：____','变量 _ 未定义','见下图__','A_','成本^收益'):
            self.assertEqual(parse(text).nodes[0].text,text)
        self.assertEqual(parse('A_').notices[0].code, 'MISSING_SCRIPT')

    def test_domain_consumption(self):
        register_domain('probe_review',lambda op,arg:('bar',1) if op=='^' and arg.startswith('c') else None)
        nodes=parse('A^cB^c','probe_review').nodes
        self.assertEqual([n.kind for n in nodes],['bar','bar'])
        self.assertEqual([n.children[0].text for n in nodes],['A','B'])
        self.assertEqual(parse('A^{c}','probe_review').nodes[-1].kind,'bar')
        self.assertEqual(parse('A^C').nodes[-1].kind,'bar')

    def test_rule_consulted_once(self):
        seen=[]
        def length_sensitive(op,arg):
            seen.append((op,arg))
            return ('bar',1) if op=='^' and len(arg)>1 and arg.startswith('c') else None
        register_domain('once_review',length_sensitive)
        self.assertEqual([n.kind for n in parse('A^cB','once_review').nodes],['bar','text'])
        self.assertEqual(list(b for b in seen if b[0]=='^'),[('^','cB')])

    def test_verifier_uppercase(self):
        path=Path(__file__).resolve().parents[1]/'scripts'/'verify_docx.py'
        spec=importlib.util.spec_from_file_location('verify_uppercase',path);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as tmp:
            for text in ('A^C','A^{C }'):
                file=Path(tmp)/'bad.docx';d=Document();d.add_paragraph(text);d.save(file)
                self.assertFalse(module.verify(str(file),[],min_images=0))

if __name__=='__main__':unittest.main()
