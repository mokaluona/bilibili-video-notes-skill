"""DOCX writer: explicit style, no review or acquisition side effects."""
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .domain import validate_document

@dataclass(frozen=True)
class StyleProfile:
    body_font: str = "宋体"
    heading_font: str = "黑体"
    body_size: float = 12
    table_size: float = 10
    caption_size: float = 10
    formula_size: float = 12
    title_size: float = 21
    heading_size: float = 17
    margin_cm: float = 1.7
    line_spacing: float = 1.25

    def __post_init__(self):
        for field in ("body_size", "table_size", "caption_size", "formula_size", "title_size", "heading_size", "line_spacing"):
            if not 0 < getattr(self, field) <= 72:
                raise ValueError(f"Invalid style field: {field}")
        if not 0.5 <= self.margin_cm <= 5:
            raise ValueError("margin_cm outside usable A4 range")

def asset_path(root, assets, asset_id):
    root = Path(root).resolve()
    path = (root / assets[asset_id]).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ValueError(f"Missing or escaping asset: {asset_id}")
    return path

def sample_blocks(blocks):
    """First representative of every used type, plus two images when available."""
    counts = {}
    chosen = []
    for block in blocks:
        kind = block["type"]
        limit = 2 if kind == "image" else 1
        if counts.get(kind, 0) < limit:
            chosen.append(block)
            counts[kind] = counts.get(kind, 0) + 1
    return chosen

def build(document, output, asset_root, style=None, sample=False):
    validate_document(document)
    if any(b["type"] == "formula" and b.get("format") == "latex" for b in document["blocks"]):
        raise ValueError("LaTeX rendering is not enabled yet; stage A only produces document data")
    style = style or StyleProfile()
    assets = document.get("assets", {})
    # Validate all image references even for a sample build.
    for block in document["blocks"]:
        if block["type"] == "image":
            with Image.open(asset_path(asset_root, assets, block["asset_id"])) as image:
                image.load()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(style.margin_cm)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = style.body_font, Pt(style.body_size)
    for paragraph_style in doc.styles:
        for border in paragraph_style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)

    def run(p, text, size, bold=False, font=None):
        r = p.add_run(text)
        r.font.name = font or style.body_font
        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font or style.body_font)
        r.font.size, r.bold = Pt(size), bold
        r.font.color.rgb = RGBColor(0, 0, 0)
        return r

    def paragraph(text, size=None, bold=False, paragraph_style=None):
        p = doc.add_paragraph(style=paragraph_style)
        p.paragraph_format.line_spacing = style.line_spacing
        p.paragraph_format.space_after = Pt(4)
        run(p, text, size or style.body_size, bold)
        return p

    title = paragraph(document["metadata"]["title"], style.title_size, True, "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if document["metadata"].get("source"):
        paragraph(document["metadata"]["source"], style.caption_size)
    blocks = sample_blocks(document["blocks"]) if sample else document["blocks"]
    figure = 0
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block.get("level", 1)
            p = paragraph(block["text"], style.heading_size - (level - 1), True, f"Heading {level}")
            p.paragraph_format.keep_with_next = True
            p.runs[0].font.name = style.heading_font
            p.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), style.heading_font)
        elif kind == "list":
            for item in block["items"]:
                if isinstance(item, str):
                    paragraph(item, paragraph_style="List Bullet")
                else:
                    p = doc.add_paragraph(style="List Bullet")
                    for part in item:
                        r = run(p, part["text"], style.body_size, part.get("bold", False))
                        if part.get("red"):
                            r.font.color.rgb = RGBColor(192, 0, 0)
        elif kind == "table":
            table = doc.add_table(rows=1, cols=len(block["headers"]))
            table.style = "Table Grid"
            for i, text in enumerate(block["headers"]):
                run(table.rows[0].cells[i].paragraphs[0], text, style.table_size, True)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in block["rows"]:
                cells = table.add_row().cells
                for i, text in enumerate(row):
                    run(cells[i].paragraphs[0], text, style.table_size)
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        elif kind == "image":
            figure += 1
            data = BytesIO()
            with Image.open(asset_path(asset_root, assets, block["asset_id"])) as image:
                image.convert("RGB").save(data, format="JPEG")
            data.seek(0)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(data, width=section.page_width - section.left_margin - section.right_margin)
            p = paragraph(f"图 {figure} {block['caption']}", style.caption_size)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "formula":
            p = paragraph(block["text"], style.formula_size)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            if kind in ("explanation", "tip"):
                label = paragraph("WHY 解释" if kind == "explanation" else "做题要点", bold=True)
                label.paragraph_format.keep_with_next = True
            p = paragraph(block["text"], bold=block.get("bold", False))
            if block.get("red"):
                p.runs[0].font.color.rgb = RGBColor(192, 0, 0)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return {"output": str(output), "sample": sample, "block_ids": [b["id"] for b in blocks]}
